"""Export the eligible competition members with no gender on file as a checklist
for staff (Excel: one row per member, tick M / F / Unknown). Runs on the Actions
runner under the owner connection because names are readable only there; the
file is uploaded as a private workflow artifact, never committed to the repo.
Env: DATABASE_URL, YEAR (default 2026)."""
import os, sys, psycopg2
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

DB_URL = os.environ.get("DATABASE_URL"); YEAR = int(os.environ.get("YEAR", "2026"))
if not DB_URL: sys.exit("DATABASE_URL not set")
conn = psycopg2.connect(DB_URL); cur = conn.cursor()
cur.execute("""
    select m.member_id, m.first_name, m.last_name, m.club, m.city, m.state,
           extract(year from m.birth_date)::int as birth_year,
           case when (m.membership_year - extract(year from m.birth_date)) between 16 and 18 then 'A'
                when (m.membership_year - extract(year from m.birth_date)) between 14 and 15 then 'B'
                when (m.membership_year - extract(year from m.birth_date)) between 12 and 13 then 'C' else 'D' end as grp,
           m.membership_type
    from membership.members m
    left join membership.member_gender mg on mg.member_id = m.member_id
    where mg.member_id is null
      and m.membership_type in ('Competition Athlete (17U)','Competition Athlete (AQUA Age 18+)')
      and m.birth_date is not null
      and (m.membership_year - extract(year from m.birth_date)) <= 18
      and m.membership_year = %s
    order by m.last_name, m.first_name
""", (YEAR,))
rows = cur.fetchall(); conn.close()

wb = Workbook(); ws = wb.active; ws.title = f"Unresolved {YEAR}"
hdr = ["Member ID", "First name", "Last name", "Club", "City", "State", "Birth year", "Group", "Membership type", "M", "F", "Unknown"]
ws.append(hdr)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="171F69"); c.alignment = Alignment(horizontal="center")
for r in rows:
    ws.append(list(r) + ["", "", ""])
dv = DataValidation(type="list", formula1='"X"', allow_blank=True)
ws.add_data_validation(dv)
n = len(rows) + 1
for col in ("J", "K", "L"):
    dv.add(f"{col}2:{col}{n}")
    for i in range(2, n + 1): ws[f"{col}{i}"].alignment = Alignment(horizontal="center")
    ws.column_dimensions[col].width = 10
for col, w in zip("ABCDEFGHI", (12, 16, 18, 28, 18, 7, 10, 7, 30)): ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
ws["N1"] = "Put an X in M, F or Unknown. Send the file back (or the Member ID + letter list) and it gets applied with source = staff-confirmed."
ws["N1"].font = Font(italic=True, color="6B7280")
os.makedirs("out", exist_ok=True)
wb.save(f"out/unresolved-gender-{YEAR}.xlsx")
print(f"{len(rows)} unresolved eligible members for {YEAR} written")
